"""Extract raw review strings from XLSX files or ZIP archives.

Only ``reviews!G3:G...`` is read. The output remains a JSON array of strings;
source information is retained only in memory for the duplicate-text report.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import sys
import tempfile
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from typing import BinaryIO, Iterable

import openpyxl


SHEET_NAME = "reviews"
FIRST_REVIEW_ROW = 3
KNOWN_HEADER = "ข้อความรีวิว (ไทย, anonymized)"
COMBINED_XLSX_NAME = "รวมไฟล์.xlsx"


class ExtractionError(Exception):
    """An input-selection or source-reading error that must abort the run."""


@dataclass(frozen=True)
class Source:
    zip_name: str | None
    excel_name: str
    sheet: str
    row: int

    def format(self) -> str:
        parts = []
        if self.zip_name is not None:
            parts.append(f"ZIP={self.zip_name}")
        parts.extend(
            (f"Excel={self.excel_name}", f"ชีต={self.sheet}", f"แถว={self.row}")
        )
        return " / ".join(parts)


@dataclass(frozen=True)
class Review:
    text: str
    source: Source


@dataclass
class Stats:
    zip_read: int = 0
    zip_duplicate_skipped: int = 0
    xlsx_read: int = 0
    xlsx_duplicate_skipped: int = 0
    combined_xlsx_skipped: int = 0
    seen_zip_hashes: dict[str, str] = field(default_factory=dict)
    seen_xlsx_hashes: dict[str, str] = field(default_factory=dict)


def _sha256_stream(stream: BinaryIO) -> str:
    digest = hashlib.sha256()
    for block in iter(lambda: stream.read(1024 * 1024), b""):
        digest.update(block)
    return digest.hexdigest()


def _sha256_file(path: Path) -> str:
    try:
        with path.open("rb") as stream:
            return _sha256_stream(stream)
    except OSError as exc:
        raise ExtractionError(f"อ่านไฟล์ไม่ได้: {path}: {exc}") from exc


def _xlsx_label(zip_name: str | None, excel_name: str) -> str:
    return excel_name if zip_name is None else f"{zip_name}/{excel_name}"


def _extract_records_from_workbook(
    wb, *, zip_name: str | None, excel_name: str
) -> list[Review]:
    label = _xlsx_label(zip_name, excel_name)
    if SHEET_NAME not in wb.sheetnames:
        raise ExtractionError(f"{label}: ไม่พบชีตชื่อ '{SHEET_NAME}'")

    ws = wb[SHEET_NAME]
    header_value = ws["G2"].value
    header_names = {KNOWN_HEADER}
    if isinstance(header_value, str):
        header_names.add(header_value.strip())

    records: list[Review] = []
    for row_number, row in enumerate(
        ws.iter_rows(min_row=FIRST_REVIEW_ROW, min_col=7, max_col=7),
        start=FIRST_REVIEW_ROW,
    ):
        cell = row[0]
        source = Source(zip_name, excel_name, SHEET_NAME, row_number)

        # Workbooks are loaded with data_only=False so formulas remain visible.
        if cell.data_type == "f":
            print(f"  [ข้ามสูตร] {source.format()} / เซลล์=G{row_number}")
            continue

        value = cell.value
        if value is None:
            continue
        if not isinstance(value, str):
            print(
                f"  [ข้ามค่าที่ไม่ใช่ string] {source.format()} / "
                f"เซลล์=G{row_number} / ประเภท={type(value).__name__}"
            )
            continue

        stripped = value.strip()
        if not stripped:
            continue
        if stripped in header_names:
            print(f"  [ข้ามหัวตาราง] {source.format()}")
            continue

        # Preserve the exact original cell string in the output.
        records.append(Review(value, source))

    return records


def _load_records_from_bytes(
    data: bytes, *, zip_name: str | None, excel_name: str
) -> list[Review]:
    label = _xlsx_label(zip_name, excel_name)
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=False
        )
    except Exception as exc:
        raise ExtractionError(f"{label}: เปิด XLSX ไม่ได้: {exc}") from exc

    try:
        return _extract_records_from_workbook(
            wb, zip_name=zip_name, excel_name=excel_name
        )
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"{label}: อ่าน XLSX ไม่ได้: {exc}") from exc
    finally:
        wb.close()


# Retain the old helpers' list[str] interface for existing callers.
def extract_reviews_from_file(xlsx_path: Path) -> list[str]:
    try:
        data = xlsx_path.read_bytes()
        records = _load_records_from_bytes(
            data, zip_name=None, excel_name=xlsx_path.name
        )
    except ExtractionError as exc:
        if "ไม่พบชีต" in str(exc):
            raise ValueError("ไม่พบชีตชื่อ 'reviews'") from exc
        raise RuntimeError(str(exc)) from exc
    except OSError as exc:
        raise RuntimeError(f"เปิดไฟล์ไม่ได้: {exc}") from exc
    return [record.text for record in records]


def extract_reviews_from_zip_entry(
    zf: zipfile.ZipFile, entry_name: str, zip_filename: str
) -> list[str]:
    try:
        data = zf.read(entry_name)
        records = _load_records_from_bytes(
            data, zip_name=zip_filename, excel_name=entry_name
        )
    except ExtractionError as exc:
        if "ไม่พบชีต" in str(exc):
            raise ValueError("ไม่พบชีตชื่อ 'reviews'") from exc
        raise RuntimeError(str(exc)) from exc
    except Exception as exc:
        raise RuntimeError(f"อ่าน entry ไม่ได้: {exc}") from exc
    return [record.text for record in records]


def _read_xlsx_data(
    data: bytes,
    *,
    zip_name: str | None,
    excel_name: str,
    stats: Stats,
) -> list[Review]:
    label = _xlsx_label(zip_name, excel_name)
    digest = hashlib.sha256(data).hexdigest()
    first_label = stats.seen_xlsx_hashes.get(digest)
    if first_label is not None:
        stats.xlsx_duplicate_skipped += 1
        print(f"  [ข้าม XLSX ซ้ำ] {label} ซ้ำกับ {first_label}")
        return []

    records = _load_records_from_bytes(
        data, zip_name=zip_name, excel_name=excel_name
    )
    stats.seen_xlsx_hashes[digest] = label
    stats.xlsx_read += 1
    print(f"  [อ่าน XLSX] {label} -> {len(records):,} รีวิว")
    return records


def _read_xlsx_path(path: Path, stats: Stats) -> list[Review]:
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise ExtractionError(f"{path}: อ่าน XLSX ไม่ได้: {exc}") from exc
    return _read_xlsx_data(
        data, zip_name=None, excel_name=path.name, stats=stats
    )


def _xlsx_entries(zf: zipfile.ZipFile) -> list[zipfile.ZipInfo]:
    return sorted(
        (
            info
            for info in zf.infolist()
            if not info.is_dir()
            and info.filename.lower().endswith(".xlsx")
            and not PurePosixPath(info.filename).name.startswith("~$")
        ),
        key=lambda info: info.filename,
    )


def _read_zip_path(path: Path, stats: Stats) -> list[Review]:
    digest = _sha256_file(path)
    first_name = stats.seen_zip_hashes.get(digest)
    if first_name is not None:
        stats.zip_duplicate_skipped += 1
        print(f"[ข้าม ZIP ซ้ำ] {path.name} ซ้ำกับ {first_name}")
        return []

    try:
        with zipfile.ZipFile(path, "r") as zf:
            entries = _xlsx_entries(zf)
            if not entries:
                raise ExtractionError(f"{path.name}: ไม่พบ XLSX ภายใน ZIP")

            records: list[Review] = []
            for entry_info in entries:
                entry_name = entry_info.filename
                if PurePosixPath(entry_name).name == COMBINED_XLSX_NAME:
                    stats.combined_xlsx_skipped += 1
                    print(
                        f"  [ข้ามไฟล์รวม] {path.name}/{entry_name} "
                        "เพราะซ้ำกับไฟล์รายสถานที่"
                    )
                    continue
                try:
                    data = zf.read(entry_info)
                except Exception as exc:
                    raise ExtractionError(
                        f"{path.name}/{entry_name}: อ่าน XLSX จาก ZIP ไม่ได้: {exc}"
                    ) from exc
                records.extend(
                    _read_xlsx_data(
                        data,
                        zip_name=path.name,
                        excel_name=entry_name,
                        stats=stats,
                    )
                )
    except ExtractionError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise ExtractionError(f"{path.name}: เปิด ZIP ไม่ได้: {exc}") from exc

    stats.seen_zip_hashes[digest] = path.name
    stats.zip_read += 1
    return records


def _supported_files(input_dir: Path) -> tuple[list[Path], list[Path]]:
    try:
        children = list(input_dir.iterdir())
    except OSError as exc:
        raise ExtractionError(f"อ่านโฟลเดอร์ไม่ได้: {input_dir}: {exc}") from exc

    zip_files = sorted(
        (p for p in children if p.is_file() and p.suffix.lower() == ".zip"),
        key=lambda p: p.name,
    )
    xlsx_files = sorted(
        (
            p
            for p in children
            if p.is_file()
            and p.suffix.lower() == ".xlsx"
            and not p.name.startswith("~$")
        ),
        key=lambda p: p.name,
    )
    return zip_files, xlsx_files


def select_inputs(
    *, input_file: Path | None, input_dir: Path | None
) -> tuple[str, list[Path]]:
    if input_file is not None:
        if not input_file.is_file():
            raise ExtractionError(f"ไม่พบไฟล์ input: {input_file}")
        suffix = input_file.suffix.lower()
        if suffix not in {".zip", ".xlsx"}:
            raise ExtractionError("--input-file ต้องเป็นไฟล์ .zip หรือ .xlsx")
        if suffix == ".xlsx" and input_file.name.startswith("~$"):
            raise ExtractionError("ไม่อ่านไฟล์ Excel ชั่วคราวที่ขึ้นต้นด้วย ~$")
        return suffix[1:], [input_file]

    assert input_dir is not None  # argparse enforces one input option.
    if not input_dir.is_dir():
        raise ExtractionError(f"ไม่พบโฟลเดอร์ input: {input_dir}")
    zip_files, xlsx_files = _supported_files(input_dir)
    if zip_files and xlsx_files:
        raise ExtractionError(
            "พบทั้ง ZIP และ XLSX ในโฟลเดอร์เดียวกัน "
            "โปรดใช้ --input-file หรือแยกไฟล์ให้เหลือชนิดเดียว"
        )
    selected = zip_files or xlsx_files
    if not selected:
        raise ExtractionError(f"ไม่พบไฟล์ .zip หรือ .xlsx ใน: {input_dir}")

    print("ไฟล์ที่จะอ่าน (เรียงตามชื่อ):")
    for path in selected:
        print(f"  - {path.name}")
    return ("zip" if zip_files else "xlsx"), selected


def collect_records(mode: str, paths: Iterable[Path]) -> tuple[list[Review], Stats]:
    records: list[Review] = []
    stats = Stats()
    for path in paths:
        records.extend(
            _read_zip_path(path, stats)
            if mode == "zip"
            else _read_xlsx_path(path, stats)
        )
    return records, stats


# Compatibility wrappers for callers of the previous module API.
def collect_reviews_from_dir(input_dir: Path) -> list[str]:
    zip_files, paths = _supported_files(input_dir)
    if zip_files:
        raise ExtractionError("โฟลเดอร์มี ZIP; ฟังก์ชันนี้รับเฉพาะ XLSX")
    records, _ = collect_records("xlsx", paths)
    return [record.text for record in records]


def collect_reviews_from_zips(input_dir: Path) -> list[str]:
    paths, xlsx_files = _supported_files(input_dir)
    if xlsx_files:
        raise ExtractionError("โฟลเดอร์มี XLSX; ฟังก์ชันนี้รับเฉพาะ ZIP")
    records, _ = collect_records("zip", paths)
    return [record.text for record in records]


def collect_reviews(input_dir: Path) -> list[str]:
    mode, paths = select_inputs(input_file=None, input_dir=input_dir)
    records, _ = collect_records(mode, paths)
    return [record.text for record in records]


def print_file_summary(stats: Stats) -> None:
    print("\nสรุปไฟล์:")
    print(f"  ZIP ที่อ่านจริง: {stats.zip_read:,}")
    print(f"  ZIP ที่ข้ามเพราะเนื้อหาซ้ำ: {stats.zip_duplicate_skipped:,}")
    print(f"  XLSX ที่อ่านจริง: {stats.xlsx_read:,}")
    print(f"  XLSX ที่ข้ามเพราะเนื้อหาซ้ำ: {stats.xlsx_duplicate_skipped:,}")
    print(f"  XLSX รวมที่ข้าม: {stats.combined_xlsx_skipped:,}")


def print_duplicate_report(records: list[Review]) -> None:
    groups: dict[str, list[Review]] = defaultdict(list)
    first_index: dict[str, int] = {}
    for index, record in enumerate(records):
        key = record.text.strip()
        groups[key].append(record)
        first_index.setdefault(key, index)

    total = len(records)
    unique = len(groups)
    duplicate_excess = total - unique
    duplicate_groups = [items for items in groups.values() if len(items) > 1]
    duplicate_groups.sort(
        key=lambda items: (-len(items), first_index[items[0].text.strip()])
    )

    print("\nรายงานข้อความซ้ำ (เปรียบเทียบทั้งข้อความหลัง strip):")
    print(f"  ข้อความทั้งหมด: {total:,}")
    print(f"  ข้อความไม่ซ้ำ: {unique:,}")
    print(f"  รายการซ้ำส่วนเกิน: {duplicate_excess:,}")
    print(
        "  หมายเหตุ: รายงานนี้ตรวจเพียงข้อความที่เหมือนกัน "
        "ไม่ได้ยืนยันว่าเป็นรีวิวจากบุคคลเดียวกัน"
    )

    if not duplicate_groups:
        print("  ไม่พบกลุ่มข้อความซ้ำ")
        return

    print("  ตัวอย่างกลุ่มซ้ำสูงสุด (ไม่เกิน 10 กลุ่ม):")
    for number, items in enumerate(duplicate_groups[:10], start=1):
        print(f"  {number}. {items[0].text.strip()!r} — {len(items):,} ครั้ง")
        for item in items:
            print(f"     - {item.source.format()}")


def atomic_write_json(output_path: Path, texts: list[str]) -> None:
    parent = output_path.parent
    if not parent.is_dir():
        raise ExtractionError(f"ไม่พบโฟลเดอร์ปลายทาง: {parent}")

    temp_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=parent,
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as temp_file:
            temp_name = temp_file.name
            json.dump(texts, temp_file, ensure_ascii=False, indent=2)
            temp_file.write("\n")
            temp_file.flush()
            os.fsync(temp_file.fileno())
        os.replace(temp_name, output_path)
    except Exception as exc:
        if temp_name is not None:
            try:
                Path(temp_name).unlink(missing_ok=True)
            except OSError:
                pass
        raise ExtractionError(f"เขียน output ไม่สำเร็จ: {output_path}: {exc}") from exc


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="สกัด reviews!G3:G... เป็น JSON array โดยไม่แก้ข้อความต้นฉบับ"
    )
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument(
        "--input-file", metavar="FILE", help="ไฟล์ .zip หรือ .xlsx หนึ่งไฟล์"
    )
    inputs.add_argument(
        "--input-dir", metavar="DIR", help="โฟลเดอร์ที่มี ZIP หรือ XLSX ชนิดเดียว"
    )
    parser.add_argument(
        "--output", required=True, metavar="FILE", help="ไฟล์ JSON ปลายทาง"
    )
    return parser.parse_args(argv)


def main() -> None:
    args = parse_args()
    input_file = Path(args.input_file) if args.input_file else None
    input_dir = Path(args.input_dir) if args.input_dir else None
    output_path = Path(args.output)

    try:
        mode, paths = select_inputs(input_file=input_file, input_dir=input_dir)
        print(f"โหมด input: {mode.upper()}")
        print(f"ไฟล์ output: {output_path.resolve()}")

        # Any selected-source error propagates before output is replaced.
        records, stats = collect_records(mode, paths)
        print_file_summary(stats)
        print_duplicate_report(records)

        if not records:
            raise ExtractionError(
                "ไม่มีรีวิวที่สกัดได้ — ไม่แทนที่ไฟล์ output เดิม"
            )

        atomic_write_json(output_path, [record.text for record in records])
    except ExtractionError as exc:
        print(f"[ผิดพลาด] {exc}", file=sys.stderr)
        sys.exit(1)

    print(f"\nบันทึกเรียบร้อย -> {output_path}")


if __name__ == "__main__":
    main()
