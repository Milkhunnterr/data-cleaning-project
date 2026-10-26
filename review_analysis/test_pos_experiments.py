import csv
import json
import math
from collections import Counter

import pytest
from openpyxl import load_workbook
from pythainlp.corpus import thai_stopwords
from pythainlp.tag import orchid

from review_analysis import analyze_tfidf as tf
from review_analysis import cluster_terms as cl
from review_analysis import pos_rules as pos
from review_analysis import run_pos_experiments as runner


@pytest.mark.parametrize("tag,expected", [
    ("NOUN", (True, True, True)), ("PROPN", (True, True, True)),
    ("VERB", (True, True, True)), ("AUX", (False, True, False)),
    ("ADJ", (False, False, True)), ("ADV", (False, False, False)),
    ("CCONJ", (False, False, False)), ("SCONJ", (False, False, False)),
    ("PRON", (False, True, False)), ("X", (False, True, False)),
    ("UNKNOWN", (False, True, False)), (None, (False, True, False)),
])
def test_policy_matrix(tag, expected):
    assert pos.keep_pos("baseline", tag)
    assert tuple(pos.keep_pos(m, tag) for m in pos.MODES[1:]) == expected


def test_library_mapping_vatt_is_adjective():
    assert orchid.TO_UD["VATT"] == "ADJ"
    assert orchid.TO_UD["VACT"] == orchid.TO_UD["VSTA"] == "VERB"
    assert orchid.TO_UD["NPRP"] == "PROPN"
    assert orchid.TO_UD["JCRG"] == "CCONJ"
    assert orchid.TO_UD["JSBR"] == "SCONJ"


def test_original_boundaries_and_no_pos_bridging():
    tokens = ["หาด", " ", "สวย", " ", "ทะเล", "\n", "เรือ", "[URL]", "เกาะ", ".", "น้ำ"]
    tags = ["NOUN", "PUNCT", "ADJ", "PUNCT", "NOUN", "PUNCT", "NOUN", "NOUN", "NOUN", "PUNCT", "NOUN"]
    filtered = Counter()
    counts, _ = tf.extract_features(tokens, set(), mode="noun_verb", tags=tags, filtered_pos=filtered)
    assert not any(f.kind == "bigram" for f in counts)
    assert filtered == {"ADJ": 1}
    counts, _ = tf.extract_features(tokens, set(), mode="noun_verb_adj", tags=tags)
    assert counts[tf.Feature("bigram", ("หาด", "สวย"))] == 1
    assert counts[tf.Feature("bigram", ("สวย", "ทะเล"))] == 1


def test_context_occurrences_df_examples_and_original_N():
    reviews = ["ดี ดี", "ดี", "ดี มาก"]
    original = list(reviews)
    documents = [["ดี", " ", "ดี"], ["ดี"], ["ดี", " ", "มาก"]]
    tags = [["VERB", "PUNCT", "ADJ"], ["ADJ"], ["VERB", "PUNCT", "ADV"]]
    result = tf.analyze_reviews(reviews, {"มาก"}, min_df=1, mode="noun_verb", prepared_tokens=documents, prepared_tags=tags)
    assert reviews == original
    assert result.tokenized_documents == documents
    assert result.unigram_count == 1 and result.bigram_count == 0
    row = result.rows[0]
    assert row["จำนวนครั้งที่พบ"] == row["จำนวนรีวิวที่พบ"] == 2
    assert float(row["คะแนน TF-IDF รวม"]) == pytest.approx(2 * (math.log(4 / 3) + 1), abs=1e-6)
    assert row["รีวิวตัวอย่าง 3 รายการ"] == "[Input 1] ดี ดี || [Input 3] ดี มาก"
    assert result.filtered_stopwords == {"มาก": 1}
    assert result.filtered_pos == {"ADJ": 2}  # ADV stopword not counted twice.


def test_tagger_alignment_and_full_context(monkeypatch):
    tokens = ["วิว", " ", "สวย", "\n", "[PHONE]", "."]
    seen = []
    def tagger(words, **kwargs):
        seen.append((list(words), kwargs))
        return [(word, "NOUN") for word in words]
    monkeypatch.setattr(pos, "pos_tag", tagger)
    assert len(pos.tag_tokens(tokens)) == len(tokens)
    assert seen == [(tokens, {"engine": "perceptron", "corpus": "orchid_ud"})]
    monkeypatch.setattr(pos, "pos_tag", lambda words, **kw: [(word, "NOUN") for word in words[1:]])
    with pytest.raises(ValueError, match="alignment"):
        pos.tag_tokens(tokens)
    monkeypatch.setattr(pos, "pos_tag", lambda words, **kw: [(word, "NOUN") for word in reversed(words)])
    with pytest.raises(ValueError, match="alignment"):
        pos.tag_tokens(tokens)


def test_actual_tagger_preserves_special_tokens():
    tokens = tf.tokenize_review("วิว สวย\nทะเล [URL] [PHONE] 123 😊 !")
    assert len(pos.tag_tokens(tokens)) == len(tokens)


def test_invalid_alignment_and_mode():
    with pytest.raises(ValueError, match="alignment"):
        tf.extract_features(["วิว"], set(), mode="noun_verb", tags=[])
    with pytest.raises(ValueError, match="tags required"):
        tf.extract_features(["วิว"], set(), mode="noun_verb")
    with pytest.raises(ValueError, match="Unknown"):
        pos.keep_pos("invalid", "NOUN")
    with pytest.raises(ValueError, match="alignment"):
        tf.analyze_reviews(["วิว"], set(), prepared_tokens=[["ทะเล"]])


@pytest.mark.parametrize("mode", pos.MODES)
def test_exports_match_occurrence_results(tmp_path, mode):
    result = tf.analyze_reviews(["วิว สวย", "ทะเล สวย"], set(thai_stopwords()), min_df=1, mode=mode)
    tf.write_csv(tmp_path / "result.csv", result.rows)
    tf.write_excel(tmp_path / "result.xlsx", result.rows)
    with (tmp_path / "result.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    wb = load_workbook(tmp_path / "result.xlsx")
    assert wb.sheetnames == ["Unigram", "Bigram"]
    for sheet, kind in zip(wb, ("unigram", "bigram")):
        expected = [r for r in rows if r["ประเภท"] == kind]
        actual = list(sheet.iter_rows(min_row=2, values_only=True))
        assert len(actual) == len(expected)
        for values, row in zip(actual, expected):
            assert values[:6] == (int(row["อันดับ"]), row["คำ/วลี"], kind, int(row["จำนวนครั้งที่พบ"]), int(row["จำนวนรีวิวที่พบ"]), float(row["คะแนน TF-IDF รวม"]))
            assert values[6] == row["รีวิวตัวอย่าง 3 รายการ"]
            assert values[7:] == (None, None)
        assert sheet["A1"].fill.fgColor.rgb == "FF234A78"
        if actual:
            assert sheet["A2"].fill.fgColor.rgb == "FFFFFFFF"
            assert sheet["H2"].fill.fgColor.rgb == "FFFFF3D4"
    assert [int(r["อันดับ"]) for r in rows] == list(range(1, len(rows) + 1))
    wb.close()


def test_baseline_ignores_pos_and_matches_original_default(monkeypatch):
    reviews = ["วิว สวย", "วิวและทะเล", "วิว สวย\nน้ำ", "ไม่ดี"]
    stopwords = set(thai_stopwords())
    expected = tf.analyze_reviews(reviews, stopwords, min_df=1)
    def forbidden(*args):
        raise AssertionError("baseline must not tag")
    monkeypatch.setattr(tf, "tag_tokens", forbidden)
    actual = tf.analyze_reviews(reviews, stopwords, min_df=1, mode="baseline", prepared_tags=[["X"]] * len(reviews))
    assert actual == expected


def test_outputs_never_overwrite_human_work(tmp_path):
    folder = tmp_path / "baseline"
    folder.mkdir()
    protected = folder / "baseline_cluster.xlsx"
    protected.write_bytes(b"human annotation")
    paths = runner.reserve_outputs(pos.MODES, tmp_path)
    assert protected.read_bytes() == b"human annotation"
    assert all(path.parent.name.startswith("pos_run_") for path in paths.values())
    assert all(not list(path.iterdir()) for path in paths.values())


@pytest.mark.parametrize("size", [0, 1])
def test_empty_or_tiny_candidates_no_fake_groups(tmp_path, monkeypatch, size):
    def forbidden(*args):
        raise AssertionError("must not encode insufficient candidates")
    monkeypatch.setattr(cl, "encode_terms", forbidden)
    candidates = [cl.Candidate(2, "วิว", "unigram")] * size
    rows = runner.cluster_result(candidates, cl.parse_args([]))
    assert all(row.group_id == "รอตรวจสอบ" for row in rows)
    assert sum(len(row.members) for row in rows) == size
    path = tmp_path / "clusters.xlsx"
    cl.write_excel(path, rows)
    cl.validate_written_workbook(path, candidates, rows)


def test_empty_pipeline_creates_all_three_outputs(tmp_path, monkeypatch, capsys):
    source = tmp_path / "input.json"
    source.write_text(json.dumps(["และ", "ไม่"], ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(tf, "OUTPUTS_DIR", tmp_path / "outputs")
    assert runner.main(["--input", str(source), "--mode", "all", "--token-examples", "0"]) == 0
    assert "N/A" in capsys.readouterr().out
    for mode in pos.MODES:
        assert sorted(p.name for p in (tmp_path / "outputs" / mode).iterdir()) == [f"{mode}_cluster.xlsx", f"{mode}_tfidf.csv", f"{mode}_tfidf.xlsx"]
