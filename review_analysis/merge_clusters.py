"""Merge intact first-pass clusters using normalized member centroids."""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
import hashlib
from io import BytesIO
import math
from pathlib import Path
import re
import sys
from zipfile import BadZipFile

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.cluster import AgglomerativeClustering

from review_analysis import cluster_terms as first
from review_analysis.pos_rules import MODES

OUTPUTS_DIR = first.BASE_DIR / "outputs"
DEFAULT_DISTANCE_THRESHOLD = 0.25  # Experimental, not a validated optimum.
SHEET_NAME = "กลุ่มคำ"
COLUMNS = ["รหัสกลุ่มใหม่", "รหัสกลุ่มเดิมที่นำมารวม", "จำนวนกลุ่มเดิม",
           "คำ/วลีทั้งหมดในกลุ่ม", "จำนวนคำ/วลี", "ชื่อแง่มุมที่ผู้ตรวจยืนยัน", "หมายเหตุผู้ตรวจ"]
WIDTHS = (18, 28, 18, 88, 18, 32, 32)


@dataclass(frozen=True)
class SourceGroup:
    group_id: str
    members: tuple[str, ...]


@dataclass(frozen=True)
class MergedGroup:
    group_id: str
    source_ids: tuple[str, ...]
    members: tuple[str, ...]


def distance_threshold(value):
    number = float(value)
    if not math.isfinite(number) or not 0 <= number <= 2:
        raise argparse.ArgumentTypeError("cosine distance threshold ต้องอยู่ระหว่าง 0 ถึง 2")
    return number


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=("all", *MODES), default="all")
    parser.add_argument("--distance-threshold", type=distance_threshold, default=DEFAULT_DISTANCE_THRESHOLD)
    parser.add_argument("--device", default="auto")
    return parser.parse_args(argv)


def checksum(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def assert_unchanged(path, expected):
    if checksum(path) != expected:
        raise ValueError(f"Input checksum เปลี่ยนระหว่างรัน: {path}")


def check_excel_lock(path):
    path = Path(path)
    if path.with_name("~$" + path.name).exists():
        raise PermissionError(f"พบไฟล์ล็อก Excel: {path}; กรุณาปิด Excel ก่อนรัน")


def load_groups(path):
    """Read only A/B/C; reviewer annotations never enter embeddings."""
    check_excel_lock(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != [first.SHEET_NAME]:
            raise ValueError(f"Input ต้องมีชีต {first.SHEET_NAME!r} เพียงชีตเดียว: {path}")
        sheet = workbook[first.SHEET_NAME]
        if list(next(sheet.values)) != first.OUTPUT_COLUMNS:
            raise ValueError(f"หัวคอลัมน์ Input ไม่ตรงกับ clustering รอบแรก: {path}")
        groups, ids, members = [], set(), set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if any(cell.data_type == "f" for cell in cells[:3]):
                raise ValueError(f"Input แถว {row_number}: ห้ามสูตรในรหัส/สมาชิก/จำนวน")
            group_id, text, count = (cell.value for cell in cells[:3])
            if not isinstance(group_id, str) or not (re.fullmatch(r"G[0-9]+", group_id) or group_id == first.REVIEW_GROUP_ID):
                raise ValueError(f"Input แถว {row_number}: รหัสกลุ่มผิดรูปแบบ {group_id!r}")
            if group_id in ids:
                raise ValueError(f"ID กลุ่มซ้ำ: {group_id}")
            ids.add(group_id)
            if not isinstance(text, str) or "\r" in text:
                raise ValueError(f"Input {group_id}: สมาชิกต้องเป็นข้อความคั่นด้วย newline")
            terms = tuple(text.split("\n"))
            if any(not term.strip() for term in terms):
                raise ValueError(f"Input {group_id}: สมาชิกว่าง")
            if type(count) is not int or count != len(terms):
                raise ValueError(f"Input {group_id}: จำนวนสมาชิกไม่ตรง")
            for term in terms:
                if term in members:
                    raise ValueError(f"สมาชิกซ้ำใน Input: {term!r} ({group_id})")
                members.add(term)
            groups.append(SourceGroup(group_id, terms))
        return groups
    finally:
        workbook.close()


def normalize(vectors):
    vectors = np.asarray(vectors, dtype=np.float64)
    if vectors.ndim != 2 or not np.isfinite(vectors).all():
        raise ValueError("embedding ต้องเป็นเมทริกซ์ finite")
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    if np.any(norms <= 1e-12):
        raise ValueError("embedding/centroid เป็น zero vector; ไม่สามารถใช้ cosine distance")
    return vectors / norms


def member_centroids(groups, embeddings):
    """Equal-weight mean of unit member vectors, then normalize each mean."""
    normalized = normalize(embeddings)
    if len(normalized) != sum(len(group.members) for group in groups):
        raise ValueError("จำนวน embedding ไม่ตรงกับสมาชิกกลุ่ม")
    means, offset = [], 0
    for group in groups:
        size = len(group.members)
        means.append(normalized[offset:offset + size].mean(axis=0))
        offset += size
    return normalize(means)


def merge_labels(centroids, threshold):
    threshold = distance_threshold(threshold)
    if len(centroids) < 2:
        return np.arange(len(centroids), dtype=int)
    return AgglomerativeClustering(
        n_clusters=None, metric="cosine", linkage="complete",
        distance_threshold=threshold, compute_full_tree=True,
    ).fit_predict(centroids)


def build_merged(groups, labels):
    normal = [g for g in groups if g.group_id != first.REVIEW_GROUP_ID]
    if len(normal) != len(labels):
        raise ValueError("จำนวน labels ไม่ตรงกับกลุ่มปกติ")
    buckets = {}
    for group, label in zip(normal, labels):
        if not isinstance(label, (int, np.integer)) or label < 0:
            raise ValueError("รอบรวมกลุ่มต้องมี label ปกติสำหรับทุกกลุ่มเดิม")
        buckets.setdefault(int(label), []).append(group)
    combined = [(tuple(g.group_id for g in bucket), tuple(t for g in bucket for t in g.members)) for bucket in buckets.values()]
    # Tie-break: source ID tuples lexicographically; keep input member order.
    combined.sort(key=lambda item: (-len(item[1]), item[0]))
    rows = [MergedGroup(f"M{i:03d}", ids, terms) for i, (ids, terms) in enumerate(combined, 1)]
    for group in groups:
        if group.group_id == first.REVIEW_GROUP_ID:
            rows.append(MergedGroup(group.group_id, (group.group_id,), group.members))
    validate_merge(groups, rows)
    return rows


def validate_merge(groups, rows):
    originals = {g.group_id: g.members for g in groups}
    if len(originals) != len(groups):
        raise ValueError("ID กลุ่มต้นทางซ้ำ")
    if len({r.group_id for r in rows}) != len(rows):
        raise ValueError("ID กลุ่มใหม่ซ้ำ")
    if Counter(s for r in rows for s in r.source_ids) != Counter(originals.keys()):
        raise ValueError("กลุ่มเดิมหายหรือซ้ำในการรวม")
    if Counter(t for r in rows for t in r.members) != Counter(t for g in groups for t in g.members):
        raise ValueError("สมาชิกหาย เพิ่ม หรือซ้ำในการรวม")
    for row in rows:
        expected = tuple(t for group_id in row.source_ids for t in originals[group_id])
        if row.members != expected:
            raise ValueError("สมาชิกกลุ่มเดิมถูกแยกหรือเปลี่ยนลำดับ")
        if first.REVIEW_GROUP_ID in row.source_ids:
            if row is not rows[-1] or row.source_ids != (first.REVIEW_GROUP_ID,) or row.group_id != first.REVIEW_GROUP_ID:
                raise ValueError("รอตรวจสอบต้องแยกไว้แถวสุดท้าย")


def excel_length(text):
    return len(text.encode("utf-16-le")) // 2


def style_sheet(sheet, widths, reviewer_columns=()):
    border = Border(**{side: Side(style="thin", color="FFB7B7B7") for side in ("left", "right", "top", "bottom")})
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet:
        for cell in row:
            header = cell.row == 1
            color = "FF234A78" if header else "FFFFF3D4" if cell.column in reviewer_columns else "FFFFFFFF"
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name=first.EXCEL_FONT_NAME, size=11, bold=header, color="FFFFFFFF" if header else "FF000000")
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if header else "left", vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 26 if row[0].row == 1 else min(409.5, max(30, 15 * max(str(c.value or "").count("\n") + 1 for c in row)))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def write_workbook(path, groups, rows, summary):
    validate_merge(groups, rows)
    values = [[r.group_id, "\n".join(r.source_ids), len(r.source_ids), "\n".join(r.members), len(r.members), None, None] for r in rows]
    for row in values:
        for value in row:
            if isinstance(value, str) and excel_length(value) > first.EXCEL_CELL_LIMIT:
                raise ValueError(f"กลุ่ม/รายการ {row[0]}: ข้อความเกินข้อจำกัด Excel {first.EXCEL_CELL_LIMIT:,}; หยุดโดยไม่ตัดข้อความ")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(COLUMNS)
    for row in values:
        sheet.append(row)
    for key, value in summary.items():
        print(f"{key}: {value}")
    # Preserve literal text, even members beginning with '='.
    for ws in workbook:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    style_sheet(sheet, WIDTHS, (6, 7))
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    saved = load_workbook(buffer, read_only=True, data_only=False)
    try:
        actual = list(saved[SHEET_NAME].iter_rows(min_row=2, values_only=True))
        if actual != [tuple(row) for row in values]:
            raise ValueError("Excel round-trip ไม่ตรงกับสมาชิก/รหัส/ช่องผู้ตรวจ")
    finally:
        saved.close()
    path = Path(path)
    check_excel_lock(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        path.write_bytes(buffer.getvalue())
    except PermissionError as exc:
        raise PermissionError(f"เขียน {path} ไม่ได้: ไฟล์อาจถูกล็อก กรุณาปิด Excel และตรวจสิทธิ์ก่อนรันใหม่") from exc


def run_mode(mode, threshold, device):
    source = OUTPUTS_DIR / mode / f"{mode}_cluster.xlsx"
    target = OUTPUTS_DIR / "merged_clusters" / f"{mode}_merged.xlsx"
    check_excel_lock(target)
    before = checksum(source)
    try:
        groups = load_groups(source)
        normal = [g for g in groups if g.group_id != first.REVIEW_GROUP_ID]
        config = first.parse_args([])
        first.set_reproducible_seed(config.seed)
        if normal:
            # kind is not used by encode_terms: only original member text is encoded.
            candidates = [first.Candidate(i + 2, term, "unigram") for i, term in enumerate(t for g in normal for t in g.members)]
            embeddings = first.encode_terms(candidates, config.model_name, config.model_revision, device, config.batch_size)
            labels = merge_labels(member_centroids(normal, embeddings), threshold)
        else:
            labels = []
        rows = build_merged(groups, labels)
        after = checksum(source)
        assert_unchanged(source, before)
        normal_rows = [r for r in rows if r.group_id != first.REVIEW_GROUP_ID]
        noise = sum(len(g.members) for g in groups if g.group_id == first.REVIEW_GROUP_ID)
        summary = {
            "mode": mode, "input": str(source.resolve()), "distance_threshold": threshold,
            "metric": "cosine", "linkage": "complete", "n_clusters": "None (threshold)",
            "model": config.model_name, "revision": config.model_revision,
            "member_normalization": "L2", "centroid": "unweighted mean then L2",
            "device": device, "batch_size": config.batch_size, "seed": config.seed,
            "scikit-learn": first.package_version("scikit-learn"),
            "sentence-transformers": first.package_version("sentence-transformers"),
            "numpy": first.package_version("numpy"),
            "normal_groups_before": len(normal), "normal_groups_after": len(normal_rows),
            "unmerged_original_groups": sum(len(r.source_ids) == 1 for r in normal_rows),
            "members_total": sum(len(g.members) for g in groups), "pending_members_unchanged": noise,
            "input_sha256_before": before, "input_sha256_after": after,
            "checks": "partition/members/pending preserved; input checksum unchanged",
            "caution": "Experimental threshold; fewer groups does not prove better Aspects. Human review required.",
        }
        write_workbook(target, groups, rows, summary)
        print(f"{mode}: {len(normal)} -> {len(normal_rows)} กลุ่มปกติ; รอตรวจสอบ {noise} คำ/วลี; checksum Input คงเดิม")
        print(f"Output: {target.resolve()}", flush=True)
        return rows
    finally:
        assert_unchanged(source, before)


def main(argv=None):
    args = parse_args(argv)
    modes = MODES if args.mode == "all" else (args.mode,)
    try:
        print("ปิดไฟล์ Excel ก่อนรัน; ผล merged ของโหมดที่เลือกจะถูกเขียนทับ ควร Save As ก่อนกรอกเอง")
        print(f"threshold={args.distance_threshold} (ค่าทดลอง); complete linkage บน cosine distance ของ centroid")
        device = first.resolve_device(args.device)
        for mode in modes:
            run_mode(mode, args.distance_threshold, device)
    except PermissionError as exc:
        print(f"ข้อผิดพลาด: ไฟล์ถูกล็อกหรือไม่มีสิทธิ์ กรุณาปิด Excel ก่อนรัน: {exc}", file=sys.stderr)
        return 1
    except (OSError, ValueError, BadZipFile) as exc:
        print(f"ข้อผิดพลาด: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
