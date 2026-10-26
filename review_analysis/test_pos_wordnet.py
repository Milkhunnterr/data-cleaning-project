from types import SimpleNamespace
import json
import pytest
from openpyxl import load_workbook
from review_analysis import pos_rules as p
from review_analysis import analyze_tfidf as tf
from review_analysis import run_pos_experiments as runner
from review_analysis.pos_wordnet_review import occurrence_rows, write_review, HEADERS


@pytest.mark.parametrize('tag,found,kinds,status', [
    ('NOUN', True, {'NOUN', 'VERB'}, p.CONFIRMED),
    ('PROPN', True, {'NOUN'}, p.CONFLICT),
    ('ADJ', True, {'ADV'}, p.CONFLICT),
    ('CCONJ', True, {'NOUN'}, p.CONFLICT),
    ('NOUN', False, set(), p.POS_ONLY),
    ('X', True, {'NOUN'}, p.UNKNOWN),
    (None, False, set(), p.UNKNOWN),
    ([], True, {'NOUN'}, p.UNKNOWN),
])
def test_statuses(tag, found, kinds, status):
    evidence = p.decide(tag, (found, frozenset(kinds)))
    assert evidence.status == status
    if status != p.CONFIRMED:
        assert not any(p.accepts(mode, evidence) for mode in p.MODES[1:])


def test_all_synsets_satellite_and_cache():
    calls = []
    def synsets(token, lang):
        calls.append((token, lang))
        return [SimpleNamespace(pos=lambda kind=kind: kind) for kind in ['n','n','v','a','s','r']]
    lexicon = p.WordNetLexicon(SimpleNamespace(synsets=synsets))
    assert lexicon.lookup('คำ') == (True, frozenset({'NOUN','VERB','ADJ','ADV'}))
    assert lexicon.lookup('คำ') == lexicon.lookup('คำ')
    assert calls == [('คำ', 'tha')]
    satellite = p.WordNetLexicon(SimpleNamespace(synsets=lambda *a, **kw: [SimpleNamespace(pos=lambda:'s')]))
    assert satellite.lookup('คำ')[1] == {'ADJ'}


def test_resource_failure_is_not_missing_word():
    def broken(*a, **kw):
        raise LookupError('missing corpus')
    lexicon = p.WordNetLexicon(SimpleNamespace(synsets=broken))
    with pytest.raises(ValueError, match='ทรัพยากร'):
        lexicon.lookup('คำ')
    assert not lexicon.cache


def test_context_and_unknown_barrier_and_audit(tmp_path):
    lexicon = SimpleNamespace(lookup=lambda token: (True, frozenset({'NOUN','VERB'})) if token != 'กลาง' else (False, frozenset()))
    reviews = ['คำ กลาง คำ', 'คำ']
    docs = [['คำ',' ','กลาง',' ','คำ'], ['คำ']]
    tags = [['NOUN','PUNCT','NOUN','PUNCT','VERB'], ['ADJ']]
    result = tf.analyze_reviews(reviews, set(), 1, mode='noun_verb', prepared_tokens=docs, prepared_tags=tags, lexicon=lexicon)
    assert result.unigram_count == 1 and result.bigram_count == 0
    assert result.rows[0]['จำนวนครั้งที่พบ'] == 2
    assert result.rows[0]['จำนวนรีวิวที่พบ'] == 1
    rows = list(occurrence_rows(reviews, docs, tags, set(), lexicon))
    assert [r[2] for r in rows] == [1,3,5,1]
    assert [r[6] for r in rows] == [p.CONFIRMED,p.POS_ONLY,p.CONFIRMED,p.CONFLICT]
    path = tmp_path/'pos_wordnet_review.xlsx'
    write_review(path, iter(rows), tmp_path/'cleaned_data.json')
    wb = load_workbook(path)
    assert len(wb.sheetnames) == 1
    assert list(next(wb.active.values)) == HEADERS
    assert [r[1] for r in list(wb.active.values)[1:]] == [reviews[0]]*3+[reviews[1]]
    assert wb.active.freeze_panes == 'A2'
    assert wb.active.auto_filter.ref == 'A1:K5'
    wb.close()


def test_baseline_runner_never_loads_resources(tmp_path, monkeypatch):
    def forbidden(*a, **kw):
        raise AssertionError('baseline must not load POS or WordNet')
    monkeypatch.setattr(runner, 'WordNetLexicon', forbidden)
    monkeypatch.setattr(runner, 'tag_tokens', forbidden)
    monkeypatch.setattr(tf, 'WordNetLexicon', forbidden)
    monkeypatch.setattr(tf, 'OUTPUTS_DIR', tmp_path)
    source = tmp_path/'cleaned_data.json'
    source.write_text(json.dumps(['และ']), encoding='utf8')
    audit = tmp_path/'pos_wordnet_review.xlsx'
    audit.write_bytes(b'existing audit')
    assert runner.main(['--mode','baseline','--input',str(source)]) == 0
    assert audit.read_bytes() == b'existing audit'


def test_malformed_tag_and_alignment(monkeypatch):
    monkeypatch.setattr(p, 'pos_tag', lambda *a, **kw: [('คำ',)])
    assert p.tag_tokens(['คำ']) == [None]
    monkeypatch.setattr(p, 'pos_tag', lambda *a, **kw: [('อื่น','NOUN')])
    with pytest.raises(ValueError, match='alignment'):
        p.tag_tokens(['คำ'])


def test_audit_overflow_preserves_existing_output(tmp_path):
    path = tmp_path/'audit.xlsx'
    path.write_bytes(b'old')
    row = [1, 'x'*32768, 1, 'x','NOUN','NOUN',p.CONFIRMED,True,True,True,'reason']
    with pytest.raises(ValueError, match='Excel'):
        write_review(path, [row], tmp_path/'cleaned_data.json')
    assert path.read_bytes() == b'old'
