from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
import pytest

from review_analysis import merge_clusters as merge


def make_input(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active.title = merge.first.SHEET_NAME
    workbook.active.append(merge.first.OUTPUT_COLUMNS)
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()


def test_centroids_normalize_members_before_mean():
    groups = [merge.SourceGroup("G01", ("a", "b")), merge.SourceGroup("G02", ("c",))]
    result = merge.member_centroids(groups, [[10, 0], [0, 1], [-3, 0]])
    np.testing.assert_allclose(result, [[2 ** -.5, 2 ** -.5], [-1, 0]])
    with pytest.raises(ValueError, match="zero vector"):
        merge.member_centroids(groups[:1], [[1, 0], [-1, 0]])
    with pytest.raises(ValueError, match="จำนวน"):
        merge.member_centroids(groups, [[1, 0]])


def test_real_complete_linkage_does_not_chain():
    # Adjacent angles close enough, endpoints not: single linkage would chain.
    angles = np.deg2rad([0, 30, 60, 180])
    vectors = np.column_stack((np.cos(angles), np.sin(angles)))
    labels = merge.merge_labels(vectors, .2)
    assert len(set(labels[:3])) == 2
    assert labels[3] not in labels[:3]
    assert sorted(Counter(labels).values()) == [1, 1, 2]
    assert len(set(merge.merge_labels(vectors, .01))) == 4
    assert len(set(merge.merge_labels(vectors, .6))) == 2


def test_threshold_zero_keeps_even_identical_vectors_separate():
    assert len(set(merge.merge_labels([[1, 0], [1, 0]], 0))) == 2
    assert len(set(merge.merge_labels([[1, 0], [1, 0]], .1))) == 1


def test_build_preserves_intact_groups_noise_and_tie_order():
    groups = [merge.SourceGroup("G02", ("b",)), merge.SourceGroup("G01", ("a",)),
              merge.SourceGroup("รอตรวจสอบ", ("noise",)), merge.SourceGroup("G03", ("c", "d"))]
    rows = merge.build_merged(groups, [0, 0, 1])
    assert [r.group_id for r in rows] == ["M001", "M002", "รอตรวจสอบ"]
    assert rows[0].source_ids == ("G02", "G01")
    assert rows[0].members == ("b", "a")
    assert rows[1].source_ids == ("G03",)
    assert rows[-1].members == ("noise",)
    assert merge.build_merged([], []) == []
    assert merge.merge_labels([], .25).size == 0
    assert list(merge.merge_labels([[1, 0]], .25)) == [0]
    with pytest.raises(ValueError):
        merge.validate_merge(groups, rows[:-1])
    with pytest.raises(ValueError):
        merge.validate_merge(groups, [rows[0], rows[0], *rows[1:]])


@pytest.mark.parametrize("rows,message", [
    ([("G01", "a", 1), ("G01", "b", 1)], "ID กลุ่มซ้ำ"),
    ([("G01", "a\na", 2)], "สมาชิกซ้ำ"),
    ([("G01", "a", 1), ("รอตรวจสอบ", "a", 1)], "สมาชิกซ้ำ"),
    ([("G01", "a", 2)], "จำนวนสมาชิก"),
    ([("G01", "a\n", 2)], "สมาชิกว่าง"),
    ([("oops", "a", 1)], "รหัสกลุ่ม"),
    ([("G01", 10, 1)], "ข้อความ"),
    ([("G01", "=1+1", 1)], "ห้ามสูตร"),
    ([("G01", "a", "1")], "จำนวนสมาชิก"),
])
def test_rejects_bad_input(tmp_path, rows, message):
    path = tmp_path / "bad.xlsx"
    make_input(path, rows)
    with pytest.raises(ValueError, match=message):
        merge.load_groups(path)


def test_annotations_ignored_and_bad_header_rejected(tmp_path):
    path = tmp_path / "input.xlsx"
    make_input(path, [("G01", "a\nb", 2, "human Aspect", "human note")])
    assert merge.load_groups(path) == [merge.SourceGroup("G01", ("a", "b"))]
    wb = load_workbook(path)
    wb.active["A1"] = "wrong"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="หัวคอลัมน์"):
        merge.load_groups(path)


def test_excel_style_literal_text_and_summary(tmp_path):
    groups = [merge.SourceGroup("G01", ("=literal", "a")), merge.SourceGroup("รอตรวจสอบ", ("pending",))]
    rows = merge.build_merged(groups, [0])
    path = tmp_path / "output.xlsx"
    merge.write_workbook(path, groups, rows, {"mode": "baseline", "distance_threshold": .25})
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["กลุ่มคำ"]
    sheet = wb["กลุ่มคำ"]
    assert list(next(sheet.values)) == merge.COLUMNS
    assert sheet["D2"].value == "=literal\na" and sheet["D2"].data_type == "s"
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:G3"
    for cell in sheet[1]:
        assert cell.fill.fgColor.rgb == "FF234A78" and cell.font.color.rgb == "FFFFFFFF"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            assert cell.alignment.wrap_text
            assert cell.fill.fgColor.rgb == ("FFFFF3D4" if cell.column >= 6 else "FFFFFFFF")
        assert row[5].value is None and row[6].value is None
    wb.close()


def test_oversized_cell_does_not_overwrite_old_result(tmp_path):
    groups = [merge.SourceGroup("G01", ("a" * 20000,)), merge.SourceGroup("G02", ("b" * 20000,))]
    rows = merge.build_merged(groups, [0, 0])
    path = tmp_path / "out.xlsx"
    path.write_bytes(b"previous result")
    with pytest.raises(ValueError, match="ข้อจำกัด Excel"):
        merge.write_workbook(path, groups, rows, {})
    assert path.read_bytes() == b"previous result"


def test_locked_file_reports_close_excel(tmp_path):
    path = tmp_path / "out.xlsx"
    path.with_name("~$" + path.name).touch()
    with pytest.raises(PermissionError, match="ปิด Excel"):
        merge.write_workbook(path, [], [], {})
    assert not path.exists()


@pytest.mark.parametrize("value", ["nan", "inf", "-0.1", "2.1"])
def test_invalid_threshold(value):
    with pytest.raises(SystemExit):
        merge.parse_args(["--distance-threshold", value])


def test_modes_repeat_input_checksum_and_noise_not_encoded(tmp_path, monkeypatch):
    monkeypatch.setattr(merge, "OUTPUTS_DIR", tmp_path)
    for mode in merge.MODES:
        make_input(tmp_path / mode / f"{mode}_cluster.xlsx", [
            ("G01", "a", 1, "ignored", "ignored"), ("G02", "b", 1),
            ("G03", "far", 1), ("รอตรวจสอบ", "noise", 1)])
    sources = {p: merge.checksum(p) for p in tmp_path.glob('*/*.xlsx')}
    def encode(candidates, model, revision, device, batch_size):
        assert [c.term for c in candidates] == ["a", "b", "far"]
        assert model == merge.first.DEFAULT_MODEL and revision == merge.first.DEFAULT_MODEL_REVISION
        return np.array([[1, 0], [1, .01], [-1, 0]])
    monkeypatch.setattr(merge.first, "encode_terms", encode)
    assert merge.main(["--mode", "all", "--device", "cpu"]) == 0
    outputs = tmp_path / "merged_clusters"
    before = {p: merge.checksum(p) for p in outputs.iterdir()}
    assert len(before) == 4
    assert merge.main(["--mode", "noun_verb", "--distance-threshold", "0", "--device", "cpu"]) == 0
    for p, digest in before.items():
        assert (merge.checksum(p) != digest) == (p.name == "noun_verb_merged.xlsx")
    for p, digest in sources.items():
        assert merge.checksum(p) == digest
    assert len(list(outputs.iterdir())) == 4
    wb = load_workbook(outputs / "baseline_merged.xlsx", read_only=True)
    assert wb.sheetnames == [merge.SHEET_NAME]
    assert wb[merge.SHEET_NAME].max_row == 4
    assert list(wb[merge.SHEET_NAME].values)[-1][3] == "noise"
    wb.close()


def test_detects_changed_checksum(tmp_path):
    path = tmp_path / "input"
    path.write_bytes(b"before")
    digest = merge.checksum(path)
    path.write_bytes(b"after")
    with pytest.raises(ValueError, match="checksum"):
        merge.assert_unchanged(path, digest)
