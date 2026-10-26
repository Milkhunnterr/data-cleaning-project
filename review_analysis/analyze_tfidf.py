"""Explore candidate aspect terms in cleaned Thai reviews with TF-IDF.

This script proposes unigram and bigram candidates for human review. It does
not infer aspects, cluster terms, classify sentiment, or modify the input.
"""

from __future__ import annotations

import argparse
import csv
import importlib.metadata
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

from review_analysis.pos_rules import MODES, keep_pos, tag_tokens

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - depends on local environment
    raise SystemExit(
        "ไม่พบ openpyxl: ติดตั้งด้วยคำสั่ง "
        "python -m pip install -r review_analysis/requirements-tfidf.txt"
    ) from exc

try:
    from pythainlp.corpus import thai_stopwords
    from pythainlp.tokenize import word_tokenize
except ImportError as exc:  # pragma: no cover - depends on local environment
    raise SystemExit(
        "ไม่พบ PyThaiNLP: ติดตั้งด้วยคำสั่ง python -m pip install pythainlp"
    ) from exc


BASE_DIR = Path(__file__).resolve().parent
OUTPUTS_DIR = BASE_DIR / "outputs"
DEFAULT_INPUT = OUTPUTS_DIR / "cleaned_data.json"
DEFAULT_OUTPUT = OUTPUTS_DIR / "tfidf_aspect_candidates.csv"

EXCEL_SHEETS = (("Unigram", "unigram"), ("Bigram", "bigram"))
EXCEL_COLUMN_WIDTHS = (9, 25, 12, 17, 18, 20, 88, 28, 32)
EXCEL_FONT_NAME = "Tahoma"

CSV_COLUMNS = [
    "อันดับ",
    "คำ/วลี",
    "ประเภท",
    "จำนวนครั้งที่พบ",
    "จำนวนรีวิวที่พบ",
    "คะแนน TF-IDF รวม",
    "รีวิวตัวอย่าง 3 รายการ",
    "แง่มุมที่ผู้ตรวจจัดกลุ่ม",
    "หมายเหตุผู้ตรวจ",
]

PLACEHOLDER_RE = re.compile(r"\[\s*(?:URL|PHONE)\s*\]", re.IGNORECASE)
PURE_NUMBER_RE = re.compile(r"^[0-9๐-๙]+(?:[.,][0-9๐-๙]+)*$")
FILTERED_THAI_MARKS = {"ๆ", "ฯ", "ฯลฯ"}


@dataclass(frozen=True)
class Feature:
    """A feature type and its token tuple."""

    kind: str
    tokens: tuple[str, ...]

    @property
    def text(self) -> str:
        return " ".join(self.tokens)


@dataclass
class AnalysisResult:
    rows: list[dict[str, object]]
    tokenized_documents: list[list[str]]
    filtered_stopwords: Counter[str]
    duplicate_groups: list[tuple[str, list[int]]]
    unigram_count: int
    bigram_count: int
    filtered_pos: Counter[str] = field(default_factory=Counter)


def token_decisions(tokens, stopwords, mode="baseline", tags=None):
    """Exclusive reasons: original filters first, POS only on usable tokens.

    Unknown tags fail allowlists but pass the explicit denylist. A stopword
    rejected by POS too is counted only as a stopword, never twice.
    """
    if mode not in MODES:
        raise ValueError(f"Unknown POS mode: {mode}")
    if tags is not None and len(tags) != len(tokens):
        raise ValueError("POS alignment mismatch: tag count")
    if mode != "baseline" and tags is None:
        raise ValueError("POS tags required for non-baseline mode")
    decisions = []
    for index, token in enumerate(tokens):
        reason = classify_token(token, stopwords)
        if reason == "usable" and mode != "baseline" and not keep_pos(mode, tags[index]):
            reason = "pos"
        decisions.append(reason)
    return decisions


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "เสนอคำ/วลี unigram และ bigram ด้วย TF-IDF เพื่อให้มนุษย์ตรวจและจัดกลุ่ม Aspect"
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"ไฟล์ JSON array of strings (default: {DEFAULT_INPUT})",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"พาธ CSV (default: {DEFAULT_OUTPUT}); Excel ใช้ชื่อเดียวกันนามสกุล .xlsx",
    )
    parser.add_argument(
        "--min-df",
        type=positive_int,
        default=2,
        help="จำนวนรีวิวขั้นต่ำที่ต้องพบ feature (default: 2)",
    )
    parser.add_argument(
        "--top-n",
        type=positive_int,
        default=50,
        help="จำนวนอันดับที่แสดงใน Terminal เท่านั้น (default: 50)",
    )
    parser.add_argument(
        "--sort-by",
        choices=("document_frequency", "tfidf"),
        default="document_frequency",
        help=(
            "เรียงตามจำนวนรีวิวที่พบ หรือคะแนน TF-IDF รวม "
            "(default: document_frequency)"
        ),
    )
    parser.add_argument(
        "--token-examples",
        type=nonnegative_int,
        default=3,
        help="จำนวนตัวอย่างการตัดคำจากข้อมูลจริง (default: 3)",
    )
    return parser.parse_args(argv)


def positive_int(value: str) -> int:
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("ค่าต้องไม่น้อยกว่า 1")
    return number


def nonnegative_int(value: str) -> int:
    number = int(value)
    if number < 0:
        raise argparse.ArgumentTypeError("ค่าต้องไม่น้อยกว่า 0")
    return number


def load_reviews(path: Path) -> list[str]:
    try:
        with path.open("r", encoding="utf-8-sig") as file:
            data = json.load(file)
    except FileNotFoundError as exc:
        raise ValueError(f"ไม่พบไฟล์ Input: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"Input ไม่ใช่ JSON ที่ถูกต้อง: {exc}") from exc

    if not isinstance(data, list) or not all(isinstance(item, str) for item in data):
        raise ValueError("Input ต้องเป็น JSON array ที่ทุกรายการเป็น string")
    if not data:
        raise ValueError("Input ไม่มีรีวิว")
    return data


def find_duplicate_groups(reviews: Sequence[str]) -> list[tuple[str, list[int]]]:
    """Return exact duplicate texts and their 1-based input positions."""
    positions: dict[str, list[int]] = defaultdict(list)
    for input_number, review in enumerate(reviews, start=1):
        positions[review].append(input_number)
    return [
        (text, input_numbers)
        for text, input_numbers in positions.items()
        if len(input_numbers) > 1
    ]


def is_symbol_or_punctuation(token: str) -> bool:
    visible = [char for char in token if not char.isspace()]
    return bool(visible) and all(
        unicodedata.category(char).startswith(("P", "S")) for char in visible
    )


def classify_token(token: str, stopwords: set[str]) -> str:
    """Classify a tokenizer token as usable, whitespace, or a hard boundary."""
    if not token or token.isspace():
        return "newline" if "\n" in token or "\r" in token else "space"
    if PLACEHOLDER_RE.fullmatch(token.strip()):
        return "placeholder"
    if token in stopwords:
        return "stopword"
    if PURE_NUMBER_RE.fullmatch(token):
        return "number"
    if token in FILTERED_THAI_MARKS:
        return "symbol"
    if is_symbol_or_punctuation(token):
        return "symbol"
    return "usable"


def tokenize_review(review: str) -> list[str]:
    """Tokenize with newmm while preserving whitespace and placeholders."""
    tokens: list[str] = []
    cursor = 0
    for match in PLACEHOLDER_RE.finditer(review):
        if match.start() > cursor:
            tokens.extend(
                word_tokenize(
                    review[cursor : match.start()],
                    engine="newmm",
                    keep_whitespace=True,
                )
            )
        tokens.append(match.group())
        cursor = match.end()
    if cursor < len(review):
        tokens.extend(
            word_tokenize(review[cursor:], engine="newmm", keep_whitespace=True)
        )
    return tokens


def extract_features(
    tokens: Sequence[str], stopwords: set[str], *, mode="baseline", tags=None,
    filtered_pos=None,
) -> tuple[Counter[Feature], Counter[str]]:
    """Extract features without joining across filtered tokens or boundaries.

    Ordinary spaces are transparent so ``วิว สวย`` can form a bigram. Every
    filtered non-space token and every newline resets bigram adjacency.
    """
    counts: Counter[Feature] = Counter()
    removed_stopwords: Counter[str] = Counter()
    previous_usable: str | None = None

    for index, (token, token_class) in enumerate(
        zip(tokens, token_decisions(tokens, stopwords, mode, tags))
    ):
        if token_class == "space":
            continue
        if token_class != "usable":
            if token_class == "stopword":
                removed_stopwords[token] += 1
            elif token_class == "pos" and filtered_pos is not None:
                filtered_pos[tags[index] or "UNKNOWN"] += 1
            previous_usable = None
            continue

        counts[Feature("unigram", (token,))] += 1
        if previous_usable is not None:
            counts[Feature("bigram", (previous_usable, token))] += 1
        previous_usable = token

    return counts, removed_stopwords


def choose_examples(
    feature: Feature,
    document_counts: Sequence[Counter[Feature]],
    reviews: Sequence[str],
    limit: int = 3,
) -> list[tuple[int, str]]:
    """Choose distinct original reviews that contain the computed feature."""
    examples: list[tuple[int, str]] = []
    seen_texts: set[str] = set()
    for input_number, (counts, original) in enumerate(
        zip(document_counts, reviews), start=1
    ):
        if counts[feature] and original not in seen_texts:
            examples.append((input_number, original))
            seen_texts.add(original)
            if len(examples) == limit:
                break
    return examples


def analyze_reviews(
    reviews: Sequence[str],
    stopwords: set[str],
    min_df: int = 2,
    sort_by: str = "document_frequency",
    *, mode: str = "baseline", prepared_tokens=None, prepared_tags=None,
) -> AnalysisResult:
    if min_df < 1:
        raise ValueError("min_df ต้องไม่น้อยกว่า 1")
    if sort_by not in {"document_frequency", "tfidf"}:
        raise ValueError("sort_by ต้องเป็น document_frequency หรือ tfidf")

    if mode not in MODES:
        raise ValueError(f"Unknown POS mode: {mode}")
    for prepared in (prepared_tokens, prepared_tags):
        if prepared is not None and len(prepared) != len(reviews):
            raise ValueError("Prepared document count does not match input")

    tokenized_documents: list[list[str]] = []
    document_counts: list[Counter[Feature]] = []
    filtered_stopwords: Counter[str] = Counter()
    filtered_pos: Counter[str] = Counter()
    total_counts: Counter[Feature] = Counter()
    document_frequencies: Counter[Feature] = Counter()

    for index, review in enumerate(reviews):
        tokens = tokenize_review(review) if prepared_tokens is None else prepared_tokens[index]
        if "".join(tokens) != review:
            raise ValueError(f"Token alignment mismatch: Input {index + 1}")
        tags = None
        if mode != "baseline":
            tags = tag_tokens(tokens) if prepared_tags is None else prepared_tags[index]
        counts, removed = extract_features(
            tokens, stopwords, mode=mode, tags=tags, filtered_pos=filtered_pos,
        )
        tokenized_documents.append(tokens)
        document_counts.append(counts)
        filtered_stopwords.update(removed)
        total_counts.update(counts)
        document_frequencies.update(counts.keys())

    document_total = len(reviews)
    candidates: list[tuple[Feature, int, int, float]] = []
    for feature, document_frequency in document_frequencies.items():
        if document_frequency < min_df:
            continue
        # smooth_idf=True, sublinear_tf=False, norm=None
        idf = math.log((1 + document_total) / (1 + document_frequency)) + 1
        tfidf_sum = total_counts[feature] * idf
        candidates.append(
            (feature, total_counts[feature], document_frequency, tfidf_sum)
        )

    if sort_by == "document_frequency":
        candidates.sort(
            key=lambda item: (
                -item[2],
                -item[3],
                -item[1],
                item[0].kind,
                item[0].text,
            )
        )
    else:
        candidates.sort(
            key=lambda item: (
                -item[3],
                -item[2],
                -item[1],
                item[0].kind,
                item[0].text,
            )
        )

    rows: list[dict[str, object]] = []
    for rank, (feature, total_count, document_frequency, tfidf_sum) in enumerate(
        candidates, start=1
    ):
        examples = choose_examples(feature, document_counts, reviews)
        example_text = " || ".join(
            f"[Input {input_number}] {review}"
            for input_number, review in examples
        )
        rows.append(
            {
                "อันดับ": rank,
                "คำ/วลี": feature.text,
                "ประเภท": feature.kind,
                "จำนวนครั้งที่พบ": total_count,
                "จำนวนรีวิวที่พบ": document_frequency,
                "คะแนน TF-IDF รวม": f"{tfidf_sum:.6f}",
                "รีวิวตัวอย่าง 3 รายการ": example_text,
                "แง่มุมที่ผู้ตรวจจัดกลุ่ม": "",
                "หมายเหตุผู้ตรวจ": "",
            }
        )

    return AnalysisResult(
        rows=rows,
        tokenized_documents=tokenized_documents,
        filtered_stopwords=filtered_stopwords,
        duplicate_groups=find_duplicate_groups(reviews),
        unigram_count=sum(1 for item in candidates if item[0].kind == "unigram"),
        bigram_count=sum(1 for item in candidates if item[0].kind == "bigram"),
        filtered_pos=filtered_pos,
    )


def write_csv(path: Path, rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=CSV_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def excel_path_for(csv_path: Path) -> Path:
    """Return the Excel path paired with a CSV output path."""
    return csv_path.with_suffix(".xlsx")


def _estimated_excel_row_height(values: Sequence[object]) -> float:
    """Estimate a readable wrapped row height without changing cell text."""
    line_count = 1
    for value, width in zip(values, EXCEL_COLUMN_WIDTHS):
        text = "" if value is None else str(value)
        visual_width = max(1, int(width - 2))
        wrapped_lines = sum(
            max(1, math.ceil(len(part) / visual_width))
            for part in text.splitlines() or [""]
        )
        line_count = max(line_count, wrapped_lines)
    # Excel limits row height to 409.5 points. Cell contents remain complete.
    return min(409.5, max(20.0, line_count * 15.0))


def _style_excel_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="FF234A78")
    white_fill = PatternFill("solid", fgColor="FFFFFFFF")
    reviewer_fill = PatternFill("solid", fgColor="FFFFF3D4")
    thin_gray = Side(style="thin", color="FFB7B7B7")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    base_font = Font(name=EXCEL_FONT_NAME, size=11)
    header_font = Font(
        name=EXCEL_FONT_NAME, size=11, bold=True, color="FFFFFFFF"
    )

    for column_number, width in enumerate(EXCEL_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column_number)].width = width

    sheet.row_dimensions[1].height = 26
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = border

    for row_number in range(2, sheet.max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, 10)]
        sheet.row_dimensions[row_number].height = _estimated_excel_row_height(values)
        for column_number in range(1, 10):
            cell = sheet.cell(row_number, column_number)
            cell.font = base_font
            cell.border = border
            cell.fill = reviewer_fill if column_number in (8, 9) else white_fill
            if column_number == 7:
                cell.alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=True
                )
            else:
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )

        for column_number in (1, 4, 5):
            sheet.cell(row_number, column_number).number_format = "0"
        sheet.cell(row_number, 6).number_format = "0.000000"

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:I{sheet.max_row}"


def write_excel(path: Path, rows: Sequence[dict[str, object]]) -> None:
    """Write the ranked result rows to separate unigram and bigram sheets."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default_sheet = workbook.active

    for sheet_index, (sheet_name, feature_type) in enumerate(EXCEL_SHEETS):
        sheet = default_sheet if sheet_index == 0 else workbook.create_sheet()
        sheet.title = sheet_name
        sheet.append(CSV_COLUMNS)

        for row in rows:
            if row["ประเภท"] != feature_type:
                continue
            values = [row[column] for column in CSV_COLUMNS]
            values[0] = int(values[0])
            values[3] = int(values[3])
            values[4] = int(values[4])
            values[5] = float(values[5])
            sheet.append(values)

        _style_excel_sheet(sheet)

    workbook.save(path)


def compact_text(text: str, limit: int = 90) -> str:
    one_line = " ".join(text.split())
    return one_line if len(one_line) <= limit else one_line[: limit - 1] + "…"


def print_duplicate_report(
    review_count: int, duplicate_groups: Sequence[tuple[str, list[int]]]
) -> None:
    duplicate_entries = sum(len(indexes) - 1 for _, indexes in duplicate_groups)
    print("\n=== ตรวจข้อความซ้ำก่อนวิเคราะห์ (เทียบข้อความตรงกันทุกตัวอักษร) ===")
    print(f"รีวิวทั้งหมด: {review_count:,}")
    print(f"ข้อความไม่ซ้ำ: {review_count - duplicate_entries:,}")
    print(f"รายการซ้ำส่วนเกิน: {duplicate_entries:,}")
    print(f"กลุ่มข้อความซ้ำ: {len(duplicate_groups):,}")
    if duplicate_groups:
        print("ตำแหน่ง Input ของทุกกลุ่มที่ซ้ำ (เริ่มนับจาก 1):")
        for text, positions in duplicate_groups:
            joined = ", ".join(map(str, positions))
            print(f"  Input {joined}: {compact_text(text)}")
    print("ไม่มีการลบหรือเปลี่ยนรายการซ้ำในการคำนวณ")


def print_token_examples(
    reviews: Sequence[str], tokenized_documents: Sequence[Sequence[str]], count: int
) -> None:
    print("\n=== ตัวอย่างการตัดคำจากข้อมูลจริง: PyThaiNLP engine=newmm ===")
    if count == 0:
        print("ปิดการแสดงตัวอย่าง (--token-examples 0)")
        return
    for index, (review, tokens) in enumerate(
        zip(reviews, tokenized_documents), start=1
    ):
        if index > count:
            break
        print(f"[Input {index}] {review}")
        print(f"Tokens: {tokens}")


def package_version(name: str) -> str:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return "ไม่ทราบ"


def print_stopword_report(removed: Counter[str], stopwords: set[str]) -> None:
    print("\n=== Stopwords จาก PyThaiNLP ===")
    print(f"PyThaiNLP version: {package_version('pythainlp')}")
    print(f"จำนวน stopwords ทั้งหมด: {len(stopwords):,}")
    print("ตรวจคำสำคัญในรายการ library:")
    for word in ("ไม่", "ราคา", "วิว", "บริการ", "สะอาด"):
        print(f"  {word}: {'อยู่' if word in stopwords else 'ไม่อยู่'}")
    print("คำที่พบและคัดออกจากสำเนาสำหรับวิเคราะห์:")
    if not removed:
        print("ไม่พบ stopword จากรายการของ PyThaiNLP")
        return
    print(", ".join(f"{word} ({count:,})" for word, count in sorted(removed.items())))


def print_top_rows(rows: Sequence[dict[str, object]], top_n: int) -> None:
    print(f"\n=== ผลอันดับต้น {min(top_n, len(rows)):,} รายการ ===")
    headers = ("อันดับ", "คำ/วลี", "ประเภท", "จำนวนครั้งที่พบ", "จำนวนรีวิวที่พบ", "คะแนน TF-IDF รวม")
    print("\t".join(headers))
    for row in rows[:top_n]:
        print("\t".join(str(row[header]) for header in headers))


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        excel_output = excel_path_for(args.output)
        if excel_output.resolve() == args.output.resolve():
            raise ValueError("--output ต้องเป็นพาธ CSV ไม่ใช่ไฟล์ .xlsx")
        if (
            args.output.resolve() == args.input.resolve()
            or excel_output.resolve() == args.input.resolve()
        ):
            raise ValueError("ไฟล์ Output ต้องไม่ทับไฟล์ Input")
        reviews = load_reviews(args.input)
        stopwords = set(thai_stopwords())
        result = analyze_reviews(
            reviews,
            stopwords,
            min_df=args.min_df,
            sort_by=args.sort_by,
        )
        print_duplicate_report(len(reviews), result.duplicate_groups)
        print_token_examples(reviews, result.tokenized_documents, args.token_examples)
        print_stopword_report(result.filtered_stopwords, stopwords)
        print("\n=== พารามิเตอร์ ===")
        print("หน่วยเอกสาร: 1 รีวิว")
        print("TF: จำนวนครั้งดิบ (sublinear_tf=False)")
        print("IDF: log((1 + N) / (1 + df)) + 1 (smooth_idf=True)")
        print("norm=None: ไม่ปรับตามความยาว รีวิวที่ยาวหรือกล่าวคำซ้ำอาจมีอิทธิพลมากขึ้น")
        print(f"min_df={args.min_df}, sort_by={args.sort_by}, top_n={args.top_n}")
        print_top_rows(result.rows, args.top_n)
        write_csv(args.output, result.rows)
        write_excel(excel_output, result.rows)
    except (OSError, ValueError) as exc:
        print(f"ข้อผิดพลาด: {exc}", file=sys.stderr)
        return 1

    print("\n=== ส่งออก ===")
    print(
        f"CSV และ Excel ทุก feature ที่ผ่าน min_df: {len(result.rows):,} รายการ "
        f"(unigram {result.unigram_count:,}, bigram {result.bigram_count:,})"
    )
    print(f"CSV: {args.output.resolve()}")
    print(f"Excel: {excel_output.resolve()}")
    print("CSV encoding: UTF-8 BOM; ช่อง Aspect และหมายเหตุเว้นว่างให้ผู้ตรวจกรอก")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
