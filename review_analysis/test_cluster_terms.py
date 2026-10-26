import csv
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from review_analysis.cluster_terms import (
    EXCEL_CELL_LIMIT,
    Candidate,
    build_cluster_rows,
    cluster_embeddings,
    load_candidates,
    parse_args,
    validate_written_workbook,
    write_excel,
)


def test_default_cluster_selection_method_is_eom_and_leaf_remains_available():
    assert parse_args([]).cluster_selection_method == "eom"
    assert parse_args(["--cluster-selection-method", "leaf"]).cluster_selection_method == "leaf"


def test_load_candidates_keeps_unigram_and_bigram_identity(tmp_path: Path):
    path = tmp_path / "input.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["คำ/วลี", "ประเภท", "คะแนน TF-IDF รวม"])
        writer.writeheader()
        writer.writerows([
            {"คำ/วลี": "น้ำ", "ประเภท": "unigram", "คะแนน TF-IDF รวม": 99},
            {"คำ/วลี": "น้ำ ใส", "ประเภท": "bigram", "คะแนน TF-IDF รวม": 1},
        ])
    candidates = load_candidates(path)
    assert [(x.term, x.kind) for x in candidates] == [("น้ำ", "unigram"), ("น้ำ ใส", "bigram")]


def test_build_rows_sorts_groups_and_places_review_last():
    candidates = [Candidate(i + 2, f"คำ{i}", "unigram") for i in range(6)]
    rows = build_cluster_rows(candidates, [8, 7, 8, -1, 8, 7])
    assert [row.group_id for row in rows] == ["G01", "G02", "รอตรวจสอบ"]
    assert [len(row.members) for row in rows] == [3, 2, 1]
    assert rows[0].member_text == "คำ0\nคำ2\nคำ4"


def test_hdbscan_can_leave_noise_instead_of_forcing_every_term():
    embeddings = np.asarray([[1, 0], [0.99, 0.01], [0.98, 0.02], [-1, 0], [-0.99, -0.01], [0, 1]], dtype=np.float32)
    embeddings /= np.linalg.norm(embeddings, axis=1, keepdims=True)
    labels, _ = cluster_embeddings(embeddings, min_cluster_size=3, min_samples=2, cluster_selection_epsilon=0, cluster_selection_method="eom", min_membership_probability=0)
    assert -1 in labels


def test_excel_has_exact_layout_and_counts(tmp_path: Path):
    candidates = [Candidate(2, "ทะเล", "unigram"), Candidate(3, "วิว สวย", "bigram"), Candidate(4, "น้ำ", "unigram")]
    rows = build_cluster_rows(candidates, [0, 0, -1])
    output = tmp_path / "clusters.xlsx"
    write_excel(output, rows)
    validate_written_workbook(output, candidates, rows)
    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["กลุ่มคำ"]
    sheet = workbook["กลุ่มคำ"]
    assert sheet.max_column == 5
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:E{sheet.max_row}"
    assert sheet["A1"].fill.fgColor.rgb == "FF234A78"
    assert sheet["D2"].fill.fgColor.rgb == "FFFFF3D4"
    assert sheet["E2"].fill.fgColor.rgb == "FFFFF3D4"
    assert sheet["D2"].value is None and sheet["E2"].value is None


def test_refuses_to_truncate_an_oversized_group():
    candidate = Candidate(2, "ก" * (EXCEL_CELL_LIMIT + 1), "unigram")
    with pytest.raises(ValueError, match="เกินข้อจำกัดเซลล์ Excel"):
        build_cluster_rows([candidate], [0])
