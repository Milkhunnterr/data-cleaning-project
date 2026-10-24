import csv
import math
from pathlib import Path

import pytest
from openpyxl import load_workbook

from review_analysis.analyze_tfidf import (
    Feature,
    analyze_reviews,
    extract_features,
    tokenize_review,
    write_csv,
    write_excel,
)


def feature_count(text, feature, stopwords=None):
    tokens = tokenize_review(text)
    counts, _ = extract_features(tokens, stopwords or set())
    return counts[feature]


def test_bigram_allows_ordinary_space():
    assert feature_count("วิว สวย", Feature("bigram", ("วิว", "สวย"))) == 1


@pytest.mark.parametrize(
    "text,stopwords",
    [
        ("วิว.สวย", set()),
        ("วิว! สวย", set()),
        ("วิว\nสวย", set()),
        ("วิว😊สวย", set()),
        ("วิวๆสวย", set()),
        ("วิวฯสวย", set()),
        ("วิว 123 สวย", set()),
        ("วิว [URL] สวย", set()),
        ("วิวและสวย", {"และ"}),
    ],
)
def test_bigram_does_not_cross_filtered_token_or_boundary(text, stopwords):
    feature = Feature("bigram", ("วิว", "สวย"))
    assert feature_count(text, feature, stopwords) == 0


def test_unigram_and_bigram_counts_document_frequency_and_tfidf():
    reviews = ["วิว สวย วิว", "วิว สวย", "ทะเล"]
    result = analyze_reviews(reviews, set(), min_df=1)
    by_key = {(row["ประเภท"], row["คำ/วลี"]): row for row in result.rows}

    view = by_key[("unigram", "วิว")]
    assert view["จำนวนครั้งที่พบ"] == 3
    assert view["จำนวนรีวิวที่พบ"] == 2
    expected = 3 * (math.log((1 + 3) / (1 + 2)) + 1)
    assert float(view["คะแนน TF-IDF รวม"]) == pytest.approx(expected, abs=1e-6)

    phrase = by_key[("bigram", "วิว สวย")]
    assert phrase["จำนวนครั้งที่พบ"] == 2
    assert phrase["จำนวนรีวิวที่พบ"] == 2


def test_default_sort_uses_df_then_tfidf_and_examples_match_features():
    reviews = ["วิว วิว", "ทะเล", "ทะเล", "วิว สวย", "วิว สวย"]
    result = analyze_reviews(reviews, set(), min_df=2)

    assert result.rows[0]["คำ/วลี"] == "วิว"
    phrase = next(row for row in result.rows if row["คำ/วลี"] == "วิว สวย")
    assert "[Input 4] วิว สวย" in phrase["รีวิวตัวอย่าง 3 รายการ"]
    assert "[Input 5] วิว สวย" not in phrase["รีวิวตัวอย่าง 3 รายการ"]
    assert "ทะเล" not in phrase["รีวิวตัวอย่าง 3 รายการ"]
    assert phrase["แง่มุมที่ผู้ตรวจจัดกลุ่ม"] == ""
    assert phrase["หมายเหตุผู้ตรวจ"] == ""


def test_tfidf_sort_option_uses_tfidf_then_document_frequency():
    reviews = ["วิว วิว วิว วิว", "ทะเล", "ทะเล", "ทะเล"]
    result = analyze_reviews(reviews, set(), min_df=1, sort_by="tfidf")
    assert result.rows[0]["คำ/วลี"] == "วิว"


def test_duplicate_reviews_are_reported_but_not_removed():
    result = analyze_reviews(["วิว", "วิว"], set(), min_df=1)
    row = next(row for row in result.rows if row["คำ/วลี"] == "วิว")
    assert row["จำนวนครั้งที่พบ"] == 2
    assert row["จำนวนรีวิวที่พบ"] == 2
    assert result.duplicate_groups == [("วิว", [1, 2])]


def test_csv_has_bom_and_all_rows(tmp_path: Path):
    result = analyze_reviews(["วิว สวย", "ทะเล สวย"], set(), min_df=1)
    output = tmp_path / "result.csv"
    write_csv(output, result.rows)

    assert output.read_bytes().startswith(b"\xef\xbb\xbf")
    with output.open(encoding="utf-8-sig", newline="") as file:
        exported = list(csv.DictReader(file))
    assert len(exported) == len(result.rows)


def test_excel_matches_csv_and_uses_required_sheet_layout(tmp_path: Path):
    result = analyze_reviews(
        ["วิว สวย", "ทะเล สวย", "วิวดี"], set(), min_df=1
    )
    csv_output = tmp_path / "result.csv"
    excel_output = tmp_path / "result.xlsx"
    write_csv(csv_output, result.rows)
    write_excel(excel_output, result.rows)

    with csv_output.open(encoding="utf-8-sig", newline="") as file:
        csv_rows = list(csv.DictReader(file))
    workbook = load_workbook(excel_output, data_only=True)

    assert workbook.sheetnames == ["Unigram", "Bigram"]
    excel_row_count = sum(sheet.max_row - 1 for sheet in workbook.worksheets)
    assert excel_row_count == len(csv_rows)

    for sheet_name, expected_type in (
        ("Unigram", "unigram"),
        ("Bigram", "bigram"),
    ):
        sheet = workbook[sheet_name]
        assert sheet.freeze_panes == "C2"
        assert sheet.auto_filter.ref == f"A1:I{sheet.max_row}"
        assert sheet["A1"].fill.fgColor.rgb == "FF234A78"
        assert sheet["H2"].fill.fgColor.rgb == "FFFFF3D4"
        assert sheet["I2"].fill.fgColor.rgb == "FFFFF3D4"

        expected_rows = [
            row for row in csv_rows if row["ประเภท"] == expected_type
        ]
        actual_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert len(actual_rows) == len(expected_rows)

        for actual, expected in zip(actual_rows, expected_rows):
            assert actual[0] == int(expected["อันดับ"])
            assert actual[1] == expected["คำ/วลี"]
            assert actual[2] == expected_type
            assert actual[3] == int(expected["จำนวนครั้งที่พบ"])
            assert actual[4] == int(expected["จำนวนรีวิวที่พบ"])
            assert actual[5] == pytest.approx(
                float(expected["คะแนน TF-IDF รวม"])
            )
            assert actual[6] == expected["รีวิวตัวอย่าง 3 รายการ"]
            assert actual[7] is None
            assert actual[8] is None


def test_thai_repetition_mark_is_not_a_feature():
    result = analyze_reviews(["สวยๆ สวย"], set(), min_df=1)
    assert all(row["คำ/วลี"] != "ๆ" for row in result.rows)
