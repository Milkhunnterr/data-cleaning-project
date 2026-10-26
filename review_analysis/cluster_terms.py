"""Cluster TF-IDF candidate terms with multilingual semantic embeddings.

The unit clustered is one original unigram/bigram candidate row, not a review
and not its numeric TF-IDF score. Uncertain terms remain available for review.
"""

from __future__ import annotations

import argparse
import csv
import importlib.metadata
import math
import random
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.cluster import HDBSCAN

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = BASE_DIR / "outputs" / "tfidf_aspect_candidates.csv"
DEFAULT_OUTPUT = BASE_DIR / "outputs" / "term_clusters.xlsx"
DEFAULT_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
DEFAULT_MODEL_REVISION = "e8f8c211226b894fcb81acc59f3b34ba3efd5f42"
REQUIRED_INPUT_COLUMNS = {"คำ/วลี", "ประเภท"}
OUTPUT_COLUMNS = ["รหัสกลุ่ม", "คำ/วลีในกลุ่ม", "จำนวนคำ/วลี", "ชื่อแง่มุมที่ผู้ตรวจยืนยัน", "หมายเหตุผู้ตรวจ"]
SHEET_NAME = "กลุ่มคำ"
REVIEW_GROUP_ID = "รอตรวจสอบ"
EXCEL_CELL_LIMIT = 32_767
EXCEL_FONT_NAME = "Tahoma"
EXCEL_COLUMN_WIDTHS = (16, 80, 16, 32, 32)


@dataclass(frozen=True)
class Candidate:
    source_row: int
    term: str
    kind: str


@dataclass(frozen=True)
class ClusterRow:
    group_id: str
    members: tuple[Candidate, ...]

    @property
    def member_text(self) -> str:
        return "\n".join(item.term for item in self.members)


def positive_int(value: str) -> int:
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("ค่าต้องไม่น้อยกว่า 1")
    return number


def probability(value: str) -> float:
    number = float(value)
    if not 0 <= number <= 1:
        raise argparse.ArgumentTypeError("ค่าต้องอยู่ระหว่าง 0 ถึง 1")
    return number


def nonnegative_float(value: str) -> float:
    number = float(value)
    if number < 0:
        raise argparse.ArgumentTypeError("ค่าต้องไม่น้อยกว่า 0")
    return number


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="จัดกลุ่มคำ/วลี TF-IDF ด้วย multilingual embeddings และ HDBSCAN โดยไม่ตั้งชื่อ Aspect อัตโนมัติ")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--model-name", default=DEFAULT_MODEL)
    parser.add_argument("--model-revision", default=DEFAULT_MODEL_REVISION)
    parser.add_argument("--device", default="auto", help="auto, cpu, cuda หรือ device ที่ Sentence Transformers รองรับ")
    parser.add_argument("--batch-size", type=positive_int, default=64)
    parser.add_argument("--min-cluster-size", type=positive_int, default=2)
    parser.add_argument("--min-samples", type=positive_int, default=1)
    parser.add_argument("--cluster-selection-epsilon", type=nonnegative_float, default=0.0, help="รวมคลัสเตอร์ที่ห่างกันไม่เกินค่านี้บน normalized Euclidean distance")
    parser.add_argument("--cluster-selection-method", choices=("eom", "leaf"), default="eom")
    parser.add_argument("--min-membership-probability", type=probability, default=0.0, help="สมาชิกที่มีความมั่นคงต่ำกว่าค่านี้จะไปกลุ่มรอตรวจสอบ")
    parser.add_argument("--seed", type=int, default=42)
    return parser.parse_args(argv)


def load_candidates(path: Path) -> list[Candidate]:
    try:
        file = path.open("r", encoding="utf-8-sig", newline="")
    except FileNotFoundError as exc:
        raise ValueError(f"ไม่พบไฟล์ Input: {path}") from exc
    with file:
        reader = csv.DictReader(file)
        missing = REQUIRED_INPUT_COLUMNS - set(reader.fieldnames or [])
        if missing:
            raise ValueError("Input ขาดคอลัมน์ที่จำเป็น: " + ", ".join(sorted(missing)))
        candidates = []
        for source_row, row in enumerate(reader, start=2):
            term, kind = row["คำ/วลี"], row["ประเภท"]
            if term is None or not term.strip():
                raise ValueError(f"คำ/วลีว่างที่แถว CSV {source_row}")
            if "\n" in term or "\r" in term:
                raise ValueError(f"คำ/วลีที่แถว CSV {source_row} มี newline ซึ่งชนกับตัวคั่นสมาชิกใน Excel")
            if kind not in {"unigram", "bigram"}:
                raise ValueError(f"ประเภทที่แถว CSV {source_row} ต้องเป็น unigram หรือ bigram")
            candidates.append(Candidate(source_row, term, kind))
    if not candidates:
        raise ValueError("Input ไม่มีคำ/วลี")
    return candidates


def set_reproducible_seed(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    try:
        import torch
        torch.manual_seed(seed)
        if torch.cuda.is_available():
            torch.cuda.manual_seed_all(seed)
    except ImportError:
        pass


def resolve_device(requested: str) -> str:
    if requested != "auto":
        return requested
    try:
        import torch
        return "cuda" if torch.cuda.is_available() else "cpu"
    except ImportError:
        return "cpu"


def encode_terms(candidates: Sequence[Candidate], model_name: str, model_revision: str, device: str, batch_size: int) -> np.ndarray:
    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise ValueError("ไม่พบ sentence-transformers: ติดตั้งด้วย python -m pip install -r review_analysis/requirements-tfidf.txt") from exc
    model = SentenceTransformer(model_name, revision=model_revision, device=device)
    embeddings = model.encode([item.term for item in candidates], batch_size=batch_size, show_progress_bar=True, convert_to_numpy=True, normalize_embeddings=True)
    result = np.asarray(embeddings, dtype=np.float32)
    if result.ndim != 2 or result.shape[0] != len(candidates):
        raise ValueError("โมเดลคืน embedding ที่มีมิติไม่ตรงกับจำนวนคำ/วลี")
    if not np.isfinite(result).all():
        raise ValueError("พบ NaN หรือ Infinity ใน embedding")
    return result


def cluster_embeddings(embeddings: np.ndarray, *, min_cluster_size: int, min_samples: int, cluster_selection_epsilon: float, cluster_selection_method: str, min_membership_probability: float) -> tuple[np.ndarray, np.ndarray]:
    if len(embeddings) < min_cluster_size:
        return np.full(len(embeddings), -1, dtype=int), np.zeros(len(embeddings))
    # On L2-normalized vectors: squared Euclidean distance = 2(1-cosine).
    # Avoiding a dense cosine matrix keeps memory bounded for thousands of terms.
    model = HDBSCAN(min_cluster_size=min_cluster_size, min_samples=min_samples, metric="euclidean", cluster_selection_epsilon=cluster_selection_epsilon, cluster_selection_method=cluster_selection_method, n_jobs=-1)
    labels = model.fit_predict(embeddings).astype(int, copy=False)
    probabilities = np.asarray(model.probabilities_, dtype=float)
    labels = labels.copy()
    labels[probabilities < min_membership_probability] = -1
    return labels, probabilities


def build_cluster_rows(candidates: Sequence[Candidate], labels: Sequence[int]) -> list[ClusterRow]:
    if len(candidates) != len(labels):
        raise ValueError("จำนวน label ไม่ตรงกับจำนวนคำ/วลี")
    grouped: dict[int, list[Candidate]] = {}
    review: list[Candidate] = []
    for candidate, label in zip(candidates, labels):
        if int(label) < 0:
            review.append(candidate)
        else:
            grouped.setdefault(int(label), []).append(candidate)
    ordered = sorted(grouped.values(), key=lambda members: (-len(members), members[0].source_row))
    rows = [ClusterRow(f"G{index:02d}", tuple(members)) for index, members in enumerate(ordered, start=1)]
    if review:
        rows.append(ClusterRow(REVIEW_GROUP_ID, tuple(review)))
    validate_partition(candidates, rows)
    validate_excel_cell_lengths(rows)
    return rows


def validate_partition(candidates: Sequence[Candidate], rows: Sequence[ClusterRow]) -> None:
    expected = Counter((x.source_row, x.term, x.kind) for x in candidates)
    actual = Counter((x.source_row, x.term, x.kind) for row in rows for x in row.members)
    if actual != expected:
        missing, duplicated = sum((expected - actual).values()), sum((actual - expected).values())
        raise ValueError(f"ผลแบ่งกลุ่มไม่ครบถ้วน: หาย {missing:,}, เกิน/ซ้ำ {duplicated:,} รายการ")


def validate_excel_cell_lengths(rows: Sequence[ClusterRow]) -> None:
    for row in rows:
        length = len(row.member_text)
        if length > EXCEL_CELL_LIMIT:
            raise ValueError(f"สมาชิกกลุ่ม {row.group_id} ยาว {length:,} ตัวอักษร เกินข้อจำกัดเซลล์ Excel {EXCEL_CELL_LIMIT:,}; ยังไม่ได้ตัดข้อความหรือเปลี่ยนรูปแบบ กรุณาปรับพารามิเตอร์ หรืออนุมัติรูปแบบรายละเอียดหนึ่งคำ/วลีต่อแถวเพื่อเก็บข้อมูลครบ")


def _estimated_row_height(member_text: str) -> float:
    visual_width = int(EXCEL_COLUMN_WIDTHS[1] - 2)
    lines = sum(max(1, math.ceil(len(line) / visual_width)) for line in member_text.splitlines())
    return min(409.5, max(20.0, lines * 15.0))


def write_excel(path: Path, rows: Sequence[ClusterRow]) -> None:
    validate_excel_cell_lengths(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(OUTPUT_COLUMNS)
    for row in rows:
        sheet.append([row.group_id, row.member_text, len(row.members), None, None])
    header_fill = PatternFill("solid", fgColor="FF234A78")
    white_fill = PatternFill("solid", fgColor="FFFFFFFF")
    reviewer_fill = PatternFill("solid", fgColor="FFFFF3D4")
    thin_gray = Side(style="thin", color="FFB7B7B7")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    base_font = Font(name=EXCEL_FONT_NAME, size=11)
    header_font = Font(name=EXCEL_FONT_NAME, size=11, bold=True, color="FFFFFFFF")
    for column, width in enumerate(EXCEL_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 26
    for cell in sheet[1]:
        cell.fill, cell.font, cell.border = header_fill, header_font, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = _estimated_row_height(sheet.cell(row_number, 2).value)
        for column in range(1, 6):
            cell = sheet.cell(row_number, column)
            cell.font, cell.border = base_font, border
            cell.fill = reviewer_fill if column in (4, 5) else white_fill
            if column in (1, 3):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
    workbook.save(path)


def validate_written_workbook(path: Path, candidates: Sequence[Candidate], rows: Sequence[ClusterRow]) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [SHEET_NAME]:
            raise ValueError("Excel ต้องมีชีตเดียวชื่อ กลุ่มคำ")
        sheet = workbook[SHEET_NAME]
        if [cell.value for cell in sheet[1]] != OUTPUT_COLUMNS:
            raise ValueError("หัวคอลัมน์ Excel ไม่ตรงตามข้อกำหนด")
        exported_terms: list[str] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            terms = [] if not values[1] else str(values[1]).split("\n")
            if values[2] != len(terms):
                raise ValueError(f"จำนวนสมาชิกไม่ตรงในกลุ่ม {values[0]}")
            if values[3] is not None or values[4] is not None:
                raise ValueError("คอลัมน์ D และ E ต้องว่าง")
            exported_terms.extend(terms)
        expected_terms = Counter(item.term for item in candidates)
        actual_terms = Counter(exported_terms)
        if actual_terms != expected_terms:
            missing = sum((expected_terms - actual_terms).values())
            duplicated = sum((actual_terms - expected_terms).values())
            raise ValueError(f"คำ/วลีใน Excel ไม่ตรงกับ Input: หาย {missing:,}, เกิน/ซ้ำ {duplicated:,}")
        validate_partition(candidates, rows)
    finally:
        workbook.close()


def package_version(name: str) -> str:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return "ไม่ทราบ"


def print_parameters(args: argparse.Namespace, device: str) -> None:
    print("=== พารามิเตอร์ clustering ===")
    print(f"input={args.input.resolve()}\noutput={args.output.resolve()}\nmodel_name={args.model_name}\nmodel_revision={args.model_revision}")
    print(f"sentence_transformers={package_version('sentence-transformers')}, scikit_learn={package_version('scikit-learn')}, numpy={package_version('numpy')}")
    print(f"device={device}, batch_size={args.batch_size}, seed={args.seed}")
    print("embedding_normalization=L2, clustering_metric=euclidean (สัมพันธ์โดยตรงกับ cosine บนเวกเตอร์ normalized)")
    print(f"min_cluster_size={args.min_cluster_size}, min_samples={args.min_samples}, cluster_selection_epsilon={args.cluster_selection_epsilon}, cluster_selection_method={args.cluster_selection_method}, min_membership_probability={args.min_membership_probability}")


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.output.suffix.lower() != ".xlsx":
            raise ValueError("--output ต้องเป็นไฟล์ .xlsx")
        if args.input.resolve() == args.output.resolve():
            raise ValueError("ไฟล์ Output ต้องไม่ทับไฟล์ Input")
        candidates = load_candidates(args.input)
        set_reproducible_seed(args.seed)
        device = resolve_device(args.device)
        print_parameters(args, device)
        print(f"อ่าน Input {len(candidates):,} รายการ (unigram {sum(c.kind == 'unigram' for c in candidates):,}, bigram {sum(c.kind == 'bigram' for c in candidates):,})")
        embeddings = encode_terms(candidates, args.model_name, args.model_revision, device, args.batch_size)
        labels, probabilities = cluster_embeddings(embeddings, min_cluster_size=args.min_cluster_size, min_samples=args.min_samples, cluster_selection_epsilon=args.cluster_selection_epsilon, cluster_selection_method=args.cluster_selection_method, min_membership_probability=args.min_membership_probability)
        rows = build_cluster_rows(candidates, labels)
        write_excel(args.output, rows)
        validate_written_workbook(args.output, candidates, rows)
    except (OSError, ValueError) as exc:
        print(f"ข้อผิดพลาด: {exc}", file=sys.stderr)
        return 1
    normal_rows = [row for row in rows if row.group_id != REVIEW_GROUP_ID]
    review_count = sum(len(row.members) for row in rows if row.group_id == REVIEW_GROUP_ID)
    accepted_probabilities = probabilities[np.asarray(labels) >= 0]
    print("\n=== ส่งออก ===")
    print(f"กลุ่มปกติ: {len(normal_rows):,} กลุ่ม\nคำ/วลีรอตรวจสอบ: {review_count:,} รายการ")
    if len(accepted_probabilities):
        print(f"membership_probability ของสมาชิกกลุ่ม: min={accepted_probabilities.min():.3f}, median={np.median(accepted_probabilities):.3f}, max={accepted_probabilities.max():.3f}")
    print(f"ตรวจครบถ้วนแล้ว: Input = สมาชิกทุกกลุ่ม = {len(candidates):,}")
    print(f"Excel: {args.output.resolve()}")
    print("คอลัมน์ D และ E เว้นว่างให้ผู้ตรวจยืนยัน; โปรแกรมไม่ได้ตั้งชื่อ Aspect")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
